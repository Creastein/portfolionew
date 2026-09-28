from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("artifacts/cv_welli_wiselie")
DOCX_PATH = OUT_DIR / "Welli_CV_Wiselie_Junior_AI_Fullstack_Engineer.docx"


def set_cell_text(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.05):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_run_font(run, size=10, bold=False, color="000000"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_rule(paragraph, color="D9D9D9", size="6", space="1"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_heading(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=7, after=2, line=1.0)
    run = p.add_run(text.upper())
    set_run_font(run, size=10.5, bold=True, color="1F4E79")
    add_rule(p, color="B7C9DD", size="4", space="1")
    return p


def add_body_paragraph(doc, text, after=3):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=after, line=1.05)
    run = p.add_run(text)
    set_run_font(run, size=9.2)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=2.2, line=1.04)
    p.paragraph_format.left_indent = Inches(0.22 + (level * 0.18))
    p.paragraph_format.first_line_indent = Inches(-0.12)
    run = p.add_run(text)
    set_run_font(run, size=8.85)
    return p


def add_role(doc, org, location, title, dates):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4, after=0, line=1.0)
    left = p.add_run(f"{org} | {location}")
    set_run_font(left, size=9.4, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p2 = doc.add_paragraph()
    set_paragraph_spacing(p2, before=0, after=2, line=1.0)
    r1 = p2.add_run(title)
    set_run_font(r1, size=9.2, bold=True, color="333333")
    r2 = p2.add_run(f" | {dates}")
    set_run_font(r2, size=9.2, color="333333")
    return p


def add_small_label(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=1, line=1.0)
    r = p.add_run(text)
    set_run_font(r, size=8.8, bold=True, color="444444")
    return p


def set_margins(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_margins(doc)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05

    # Name block
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=0, line=1.0)
    r = p.add_run("WELLI")
    set_run_font(r, size=19, bold=True, color="0B2545")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=2, line=1.0)
    r = p.add_run("Full Stack Developer & AI Operations Engineer | Junior AI / Fullstack Engineer Candidate")
    set_run_font(r, size=9.3, bold=True, color="333333")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=4, line=1.0)
    contacts = [
        "Tangerang Selatan, Indonesia",
        "+62 851 6150 7114",
        "wellibuilds@gmail.com",
        "linkedin.com/in/welli-",
        "welli.my.id",
        "github.com/Creastein",
    ]
    r = p.add_run(" | ".join(contacts))
    set_run_font(r, size=8.3, color="333333")

    add_heading(doc, "Ringkasan Profesional")
    add_body_paragraph(
        doc,
        "Full Stack Developer & AI Operations Engineer dengan pengalaman membangun, mengelola, dan mendokumentasikan "
        "sistem AI agent, workflow otomasi, aplikasi internal berbasis web, database, dan integrasi API. Terbiasa "
        "menggunakan Next.js, React, TypeScript, Supabase/PostgreSQL, n8n, GitHub, dan deployment modern untuk "
        "mendukung kebutuhan operasional bisnis. Memiliki pengalaman langsung mengoperasikan sistem multi-agent "
        "berbasis OpenClaw dan Hermes Agent, membangun command interface melalui Telegram/Discord, serta menerjemahkan "
        "kebutuhan bisnis menjadi solusi teknologi yang stabil, terdokumentasi, aman, dan dapat dikembangkan.",
        after=4,
    )

    add_heading(doc, "Keahlian Teknis")
    skills = [
        ("Frontend", "React.js, Next.js, TypeScript, JavaScript, Tailwind CSS, HTML, CSS"),
        ("Backend & Database", "Node.js, REST API, Supabase, PostgreSQL, OAuth, Webhook, Python for automation scripting"),
        ("AI & Automation", "OpenClaw, Hermes Agent, n8n, LLM Integration, Claude API, Tool Calling, Multi-Agent Architecture, AI Workflow Design"),
        ("DevOps & Operations", "Git, GitHub, Vercel, Google Cloud Run, Docker, WSL, Cron Jobs, DNS Setup, Deployment, Server Environment, Logging"),
        ("Integrations", "Telegram Bot API, Discord Bot, WhatsApp API, Google Analytics, SerpApi, Google Workspace API concepts, Third-party API Integration"),
        ("Documentation & Ownership", "Technical SOP, Deployment Guide, System Architecture Documentation, Incident Notes, Problem Solving, Debugging"),
        ("AI Concepts", "RAG concepts, Vector Database concepts, AI Agent Operations, Automation Workflow, Monitoring Dashboard"),
    ]
    for label, value in skills:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=1.8, line=1.02)
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=8.75, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=8.75)

    add_heading(doc, "Pengalaman Kerja")
    add_role(doc, "WelliBuilds", "Tangerang Selatan, Indonesia", "AI Operations & Full Stack Developer", "2025 - Sekarang")
    add_small_label(doc, "AI Operations, Automation & Internal Systems")
    welli_ai = [
        "Merancang dan mengoperasikan sistem multi-agent AI untuk mendukung workflow operasional bisnis menggunakan OpenClaw, mencakup koordinasi CRM, sales pipeline, marketing workflow, daily briefing otomatis, dan approval-based automation melalui Telegram.",
        "Melakukan migrasi workflow AI operations dari OpenClaw ke Hermes Agent, termasuk evaluasi platform, redesign arsitektur multi-agent, integrasi ulang command interface ke Discord, dan penyesuaian workflow operasional.",
        "Mendesain pembagian role multi-agent berdasarkan fungsi bisnis, termasuk CEO agent untuk koordinasi, Sales agent untuk pipeline, Marketing agent untuk content workflow, dan AI Monitoring agent untuk pemantauan aktivitas.",
        "Membangun automation pipeline berbasis n8n, Docker, dan WSL untuk mendukung workflow konten YouTube dengan Hermes Agent via MCP sebagai orchestrator.",
        "Mengintegrasikan Telegram Bot API dan Discord Bot sebagai interface perintah, notifikasi, dan approval layer untuk sistem otomasi.",
        "Membuat dokumentasi teknis, SOP operasional, deployment guide, dan catatan arsitektur untuk sistem AI dan automation workflow.",
    ]
    for item in welli_ai:
        add_bullet(doc, item)

    add_small_label(doc, "Full Stack Development & Internal Tools")
    welli_web = [
        "Mengembangkan LeadFinder AI, internal tool berbasis Next.js 15 dan SerpApi untuk membantu proses prospecting bisnis secara otomatis.",
        "Membangun dan memelihara aplikasi web menggunakan React, Next.js, TypeScript, Tailwind CSS, REST API, dan integrasi layanan pihak ketiga.",
        "Mengelola database, autentikasi, integrasi API, dan deployment aplikasi menggunakan Supabase, PostgreSQL, GitHub, dan Vercel.",
        "Melakukan troubleshooting, debugging, deployment, DNS setup, domain configuration, dan Google Search Console untuk proyek web klien.",
        "Mengembangkan website komersial untuk sektor hospitality, termasuk thesecretkarimunjawa.com dan floatingparadise.id, dengan skor performa PageSpeed 95, SEO 100, dan Accessibility 100.",
    ]
    for item in welli_web:
        add_bullet(doc, item)

    add_role(doc, "PT. Graha Era Nusantara (Wartelsus)", "Indonesia", "Admin & Web Developer", "Sep 2025 - Des 2025")
    for item in [
        "Mendigitalisasi alur kerja internal dengan merancang dan membangun website operasional perusahaan untuk mengurangi dependensi pada pencatatan manual.",
        "Mengelola database internal dan berkoordinasi langsung dengan vendor pihak ketiga untuk kebutuhan sistem, operasional, dan integrasi data.",
    ]:
        add_bullet(doc, item)

    add_role(doc, "PT. Exi Global Indonesia", "Indonesia", "Business Analyst Intern", "Apr 2024 - Jul 2024")
    for item in [
        "Menerjemahkan kebutuhan operasional bisnis ke dalam spesifikasi teknis terstruktur untuk tim pengembang.",
        "Mendokumentasikan alur logika sistem melalui flowchart dan Entity Relationship Diagram (ERD) menggunakan draw.io.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Proyek Terpilih")
    projects = [
        ("AI Agent Operations System", "Sistem multi-agent untuk mengotomasi koordinasi CRM, sales pipeline, marketing workflow, daily briefing, dan approval-based task execution melalui Telegram/Discord."),
        ("LeadFinder AI", "Internal prospecting tool berbasis Next.js 15 dan SerpApi untuk menemukan bisnis potensial tanpa website secara otomatis."),
        ("YouTube Automation Pipeline", "Workflow otomasi berbasis n8n, Docker, WSL, dan Hermes Agent via MCP untuk mendukung proses produksi konten."),
        ("Hospitality Website Projects", "Website komersial untuk thesecretkarimunjawa.com dan floatingparadise.id dengan fokus pada performa, SEO, accessibility, deployment, dan domain setup."),
    ]
    for name, desc in projects:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=2, line=1.02)
        r = p.add_run(f"{name}: ")
        set_run_font(r, size=8.85, bold=True)
        r = p.add_run(desc)
        set_run_font(r, size=8.85)

    add_heading(doc, "Pendidikan")
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=1, line=1.0)
    r = p.add_run("Universitas Kristen Satya Wacana (UKSW) | Salatiga, Indonesia")
    set_run_font(r, size=9.1, bold=True)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=1, line=1.0)
    r = p.add_run("Sarjana Sistem Informasi | 2018 - 2025 | IPK: 3.57 / 4.00")
    set_run_font(r, size=9.1)

    # Keep the file metadata plain.
    doc.core_properties.author = "Welli"
    doc.core_properties.title = "Welli - CV - Junior AI / Fullstack Engineer"
    doc.core_properties.subject = "CV tailored for Wiselie International Solutions"
    doc.core_properties.keywords = (
        "Junior AI Engineer, Fullstack Engineer, AI Operations, OpenClaw, Supabase, PostgreSQL, "
        "React, Next.js, TypeScript, Python, Automation Workflow, API Integration"
    )

    doc.save(DOCX_PATH)
    print(DOCX_PATH.resolve())


if __name__ == "__main__":
    build()
